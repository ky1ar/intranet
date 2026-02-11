import logging, os, uuid
from werkzeug.utils import secure_filename
from flask import request, send_file, render_template
from flask_mail import Message
from datetime import  date, datetime
from application import mail
from application.handlers import handle_exceptions
from application.utils import format_name, format_datetime, calculate_passed_days, format_date
from application.repository.complaint_repository import ComplaintRepository
from application.repository.user_repository import UserRepository
from application.repository.client_repository import ClientRepository
from application.services.push_service import PushSender
from application import socketio
from flask_jwt_extended import get_jwt_identity
from config import Config


try:
    from docx import Document
except Exception:
    Document = None

try:
    import openpyxl
except Exception:
    openpyxl = None


class ComplaintService:
    def __init__(self):
        self.complaint_repository = ComplaintRepository()
        self.client_repository = ClientRepository()
        self.user_repository = UserRepository()
        self.push_service = PushSender()
        self.worker_lvl = 2
        self.leader_lvl = 3
        self.admin_lvl = 4
        self.fabrix_id = 21
        self.admin_dep = 1


    def _attachments_dir(self, path=Config.UPLOAD_PROPF_FOLDER):
        try:
            os.makedirs(path, exist_ok=True)
        except Exception:
            logging.exception("Could not create upload folder: %s", path)
            raise
        return path
    

    def _file_ext(self, filename):
        return (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()


    def _allowed_ext(self, filename):
        allowed = {
            "pdf", "xml", "txt",
            "doc", "docx",
            "xls", "xlsx",
            "png", "jpg", "jpeg", "webp", "gif"
        }
        return self._file_ext(filename) in allowed


    def _human_size(self, n):
        if n is None:
            return ""
        n = float(n)
        for unit in ["B","KB","MB","GB"]:
            if n < 1024:
                return f"{n:.0f}{unit}"
            n /= 1024
        return f"{n:.0f}TB"
    

    @handle_exceptions
    def new(self, data):
        user_id = data.get("user_id")
        send_mail = False
        client_id = data.get("client_id")
        client_data = data.pop("client")

        document = client_data.get("document", "").strip()
        name = client_data.get("name", "").strip()
        phone = client_data.get("phone", "").strip()
        email = client_data.get("email", "").strip()
        department_id = client_data.get("department_id")
        province_id = client_data.get("province_id")
        district_id = client_data.get("district_id")
        address = client_data.get("address")

        if not user_id:
            send_mail = True
            user_id = self.fabrix_id
        
        if client_id:
            if not email:
                return "Ingresa tu correo", 400
            if not department_id:
                return "Selecciona un departamento", 400
            if not province_id:
                return "Selecciona una provincia", 400
            if not district_id:
                return "Selecciona un distrito", 400
            if not address:
                return "Ingresa tu dirección", 400
            
            client, cc = self.client_repository.get_client_by_id(client_id)
            if cc != 200:
                return client, cc
            self.client_repository.update_client(client, client_data)

        else:
            if not document:
                return "Ingrese un documento", 400
            if not name:
                return "Ingrese el nombre", 400
            if not phone or len(phone) != 9:
                return "Ingresa un celular válido", 400
            if not email:
                return "Ingresa tu correo", 400
            if not department_id:
                return "Selecciona un departamento", 400
            if not province_id:
                return "Selecciona una provincia", 400
            if not district_id:
                return "Selecciona un distrito", 400
            if not address:
                return "Ingresa tu dirección", 400
            
            client, cc = self.client_repository.get_client_by_document(document)
            if cc == 500:
                return client, cc
            
            if cc == 404:
                add_client, acc = self.client_repository.add_client(client_data)
                if acc != 200:
                    return add_client, acc
                data["client_id"] = add_client
            else:
                updated_client, ucc = self.client_repository.update_client(client, client_data)
                if ucc != 200:
                    return updated_client, ucc
                
                data["client_id"] = client.id
        
        consumption_id = data.get("consumption_id")
        type_id = data.get("type_id")
        request = data.get("request")

        if not type_id:
                return "Selecciona un tipo de reclamo", 400
        if not consumption_id:
                return "Selecciona un tipo de consumo", 400
        if not request:
                return "Ingresa la solicitud del cliente", 400
        
        complaint_letter = "RC" if data.get("type_id") == 1 else "QJ"

        complaint_id, ncc = self.complaint_repository.new_complaint(data)
        if ncc != 200:
            return complaint_id, ncc

        new_history, nhc = self.complaint_repository.new_history(complaint_id, user_id, 0, "El cliente registró su reclamo")
        if nhc != 200:
            return new_history, nhc
        
        department_user_ids, duic = self.user_repository.get_user_ids_by_department(department_id=self.admin_dep)
        if duic != 200:
            return department_user_ids, duic
        
        department_user_ids.append(23)

        self.push_service.send_to_users(
            user_ids=department_user_ids,
            title=f"Nuevo Reclamo {complaint_letter}-{complaint_id}",
            body="Se ha registrado un nuevo reclamo pendiente de gestión.",
        )

        if send_mail:
            html_content = render_template(
                'claim.html',
                email=email,
                complaint_id=f"{complaint_letter}-{complaint_id}",
                current_year=datetime.now().year
            )
            msg = Message(
                subject="Confirmación de Reclamo - Krear 3D",
                sender=("Tienda Krear 3D", "web@tiendakrear3d.com"),
                recipients=[email],
                html=html_content
            )
            mail.send(msg)

        socketio.emit("complaint_dashboard_update", {})

        return "Reclamo registrado correctamente", 200
    

    @handle_exceptions
    def status(self):
        status, sc = self.complaint_repository.get_statuses()
        if sc != 200:
            return status, sc
        
        status_list = [
            {
                "id": option.id,
                "name": option.name,
                "image": option.image,

            } for option in status
        ]
        return status_list, 200
    

    @handle_exceptions
    def dashboard(self):
        user_id = int(get_jwt_identity())
        user, uc = self.user_repository.get_user_by_id(user_id)
        if uc != 200:
            return user, uc
        
        department_id = user.department_id
        if department_id in [4, 6, 8] and user_id not in [15, 23]:
            return "Not allowed", 200
        
        level_id = user.level_id
        if level_id == self.worker_lvl and department_id != 1:
            complaints, cc = self.complaint_repository.get_complaints(user_id) 
        else:
            complaints, cc = self.complaint_repository.get_complaints() 

        if cc != 200:
            return complaints, cc

        statuses, sc = self.complaint_repository.get_statuses() 
        if sc != 200:
            return statuses, sc
        
        result = []
        for status in statuses:
            complaint_status = [complaint_row for complaint_row in complaints if complaint_row.status_id == status.id]
            
            complaints_list = []
            for complaint in complaint_status:
                passed_days = 0
                updated_today = False

                first_status, _ = self.complaint_repository.get_complaint_status(complaint.id, 1) 
                if first_status:
                    passed_days = calculate_passed_days(first_status.created_at)

                current_status, _ = self.complaint_repository.get_complaint_status(complaint.id, status.id) 
                if current_status and current_status.created_at.date() == date.today():
                    updated_today = True

                complaints_data = {
                    "id": complaint.id,
                    "client_name": complaint.client.name,
                    "consumption_name": complaint.complaint_consumption.name,
                    "type_letter": "RC" if complaint.complaint_type_id == 1 else "QJ",
                    "assigned_name": format_name(complaint.owner.name, simple=True) if complaint.owner_id else None,
                    "priority": 'low',
                    "updated_today": updated_today,
                    "passed_days": passed_days,
                    "status_id": status.id,
                    "order_number": complaint.order_number,
                }
                    
                complaints_list.append(complaints_data)
            
            complaints_list.sort(key=lambda x: x["passed_days"], reverse=True)

            result.append({
                "status_id": status.id,
                "status_name": status.name,
                "complaints": complaints_list
            })
        return result, 200
    

    @handle_exceptions
    def view(self, complaint_id):
        complaint, cc = self.complaint_repository.get_complaint(complaint_id) 
        if cc != 200:
            return complaint, cc
        
        passed_days = 0
        first_status, _ = self.complaint_repository.get_complaint_status(complaint_id, 1) 
        if first_status:
            passed_days = calculate_passed_days(first_status.created_at)

        response = {
            "id": complaint_id,
            "document": complaint.client.document,
            "client_name": complaint.client.name,
            "phone": complaint.client.phone,
            "email": complaint.client.email,
            "department": complaint.client.department.name,
            "province": complaint.client.province.name,
            "district": complaint.client.district.name,
            "address": complaint.client.address,
            "priority": 'low',
            "passed_days": passed_days,

            "order_number": complaint.order_number,
            "consumption_name": complaint.complaint_consumption.name,
            "amount": complaint.amount,
            "consumption_description": complaint.consumption_description,
            "purchase_date": format_date(complaint.purchase_date),

            "detail": complaint.detail,
            "customer_request": complaint.customer_request,

            "type": complaint.complaint_type.name,
            "type_letter": "RC" if complaint.complaint_type_id == 1 else "QJ",
            "status_id": complaint.status_id,

            "owner_name": format_name(complaint.owner.name) if complaint.owner else None,
            "seller_name": format_name(complaint.seller.name) if complaint.seller else None,
            
            "chats": []
        }

        for chat in complaint.chats:
            response["chats"].append({
                "id": chat.id,
                "comment": chat.comment,
                "commenter_id": chat.commenter_id,
                "commenter_name": format_name(chat.commenter.name),
                "commenter_image": chat.commenter.image,
                "created_at": format_datetime(chat.created_at),
            })

        # history, hc = self.complaint_repository.get_complaint_history(complaint_id) 
        # if hc != 200:
        #     return history, hc
        
        # history_dict = [
        #     {
        #         "status_id": row.status_id,
        #         "user_name": format_name(row.user.name, simple=True),
        #         "notes": row.notes if row.notes else "",
        #         "created_at":  format_datetime(row.created_at)
        #     } for row in history
        # ]
        # response["history"] = history_dict
                
        return response, 200
    

    @handle_exceptions
    def move(self, data):
        user_id = int(get_jwt_identity())

        complaint_id = data.get("complaint_id")
        notes = data.get("notes")
        
        complaint, cc = self.complaint_repository.get_complaint(complaint_id) 
        if cc != 200:
            return complaint, cc

        current_status_id = complaint.status_id
        complaint_letter = "RC" if complaint.complaint_type_id == 1 else "QJ"
        
        move, mc = self.complaint_repository.move_status(complaint, current_status_id, data) 
        if mc != 200:
            return move, mc
        
        socketio.emit("complaint_dashboard_update", {})
        new_history, nhc = self.complaint_repository.new_history(complaint_id, user_id, current_status_id, notes) 
        if nhc != 200:
            return new_history, nhc

        if current_status_id == 2:
            owner_id = data.get("owner_id")
            seller_id = data.get("seller_id")

            users_to_send = []
            if owner_id:
                users_to_send.append(int(owner_id))
            if seller_id and seller_id != owner_id:
                users_to_send.append(int(seller_id))
                    
            self.push_service.send_to_users(
                user_ids=users_to_send,
                title=f"Reclamo {complaint_letter}-{complaint_id} asignado",
                body="Se te ha asignado un nuevo reclamo, por favor agrega evidencias.",
            )

        return "Reclamo actualizado correctamente", 200
    

    def attachments_list(self, complaint_id):
        complaint_id = int(complaint_id)

        ok, _ = self.complaint_repository.complaint_exists(complaint_id)
        if not ok:
            return "Complaint not found", 404

        rows, rc = self.complaint_repository.get_attachments(complaint_id)
        if rc != 200:
            return rows, rc

        data = []
        for r in rows:
            data.append({
                "id": r.id,
                "original_name": r.original_name,
                "mime_type": r.mime_type,
                "size_bytes": r.size_bytes,
                "size_h": self._human_size(r.size_bytes),
                "created_at": str(r.created_at),
                "ext": self._file_ext(r.original_name),

                "inline_url": f"/complaint/attachment/{r.id}?disposition=inline",
                "download_url": f"/complaint/attachment/{r.id}?disposition=attachment",
                "preview_url": f"/complaint/attachment/{r.id}/preview",
            })
        return data, 200


    # POST /complaint/attachments/<complaint_id>
    # form-data: files[]
    def attachments_upload(self, complaint_id):
        user_id = int(get_jwt_identity())
        complaint_id = int(complaint_id)

        ok, _ = self.complaint_repository.complaint_exists(complaint_id)
        if not ok:
            return "Complaint not found", 404

        files = request.files.getlist("files[]")
        if not files:
            return "No files received", 400

        max_bytes = Config.COMPLAINT_ATTACHMENTS_MAX_BYTES

        saved_ids = []
        for f in files:
            if not f or not f.filename:
                continue

            if not self._allowed_ext(f.filename):
                return f"File type not allowed: {f.filename}", 400

            try:
                pos = f.stream.tell()
                f.stream.seek(0, os.SEEK_END)
                size = f.stream.tell()
                f.stream.seek(pos)
            except Exception:
                size = None

            if size and size > max_bytes:
                return f"File too large (max {max_bytes} bytes): {f.filename}", 400

            original_name = f.filename
            ext = self._file_ext(original_name)
            safe = secure_filename(original_name) or f"file.{ext}"
            stored_name = f"{uuid.uuid4().hex}_{safe}"

            path = os.path.join(self._attachments_dir(), stored_name)
            f.save(path)

            mime = getattr(f, "mimetype", None)
            size_bytes = os.path.getsize(path)

            new_id, nc = self.complaint_repository.add_attachment(
                complaint_id=complaint_id,
                user_id=user_id,
                original_name=original_name,
                stored_name=stored_name,
                mime_type=mime,
                size_bytes=size_bytes,
            )
            if nc != 200:
                return new_id, nc

            saved_ids.append(new_id)

        return {"uploaded": saved_ids}, 200


    # GET /complaint/attachment/<attachment_id>?disposition=inline|attachment
    def attachment_stream(self, attachment_id, disposition="inline"):
        row, rc = self.complaint_repository.get_attachment_by_id(int(attachment_id))
        if rc != 200:
            return row, rc

        path = os.path.join(self._attachments_dir(), row.stored_name)
        if not os.path.exists(path):
            return "File missing on disk", 404

        as_attachment = (disposition == "attachment")
        return send_file(
            path,
            mimetype=row.mime_type or "application/octet-stream",
            as_attachment=as_attachment,
            download_name=row.original_name
        )


    # GET /complaint/attachment/<attachment_id>/preview
    def attachment_preview(self, attachment_id):
        row, rc = self.complaint_repository.get_attachment_by_id(int(attachment_id))
        if rc != 200:
            return row, rc

        ext = self._file_ext(row.original_name)
        inline_url = f"/complaint/attachment/{row.id}?disposition=inline"

        if ext in {"png","jpg","jpeg","webp","gif","pdf"}:
            return {"kind": "url", "url": inline_url, "name": row.original_name}, 200

        path = os.path.join(self._attachments_dir(), row.stored_name)
        if not os.path.exists(path):
            return "File missing on disk", 404

        if ext in {"txt","xml"}:
            with open(path, "r", encoding="utf-8", errors="replace") as fp:
                text = fp.read(200_000)
            return {"kind": "text", "text": text, "name": row.original_name}, 200

        if ext == "docx" and Document:
            doc = Document(path)
            parts = [p.text for p in doc.paragraphs if p.text]
            text = "\n".join(parts)[:200_000]
            return {"kind": "text", "text": text, "name": row.original_name}, 200

        if ext == "xlsx" and openpyxl:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            max_rows = min(ws.max_row or 0, 50)
            max_cols = min(ws.max_column or 0, 20)

            rows_html = []
            for r in range(1, max_rows + 1):
                cols = []
                for c in range(1, max_cols + 1):
                    v = ws.cell(row=r, column=c).value
                    cols.append(f"<td>{'' if v is None else str(v)}</td>")
                rows_html.append("<tr>" + "".join(cols) + "</tr>")

            html = "<table>" + "".join(rows_html) + "</table>"
            return {"kind": "html", "html": html, "name": row.original_name}, 200

        return {
            "kind": "download",
            "name": row.original_name,
            "message": "Vista previa no disponible",
            "download_url": f"/complaint/attachment/{row.id}?disposition=attachment"
        }, 200
    

    @handle_exceptions
    def options_type(self):
        types, tc = self.complaint_repository.get_options_type() 
        if tc != 200:
            return types, tc
        
        result = []
        for item in types:
            result.append({
                "id": item.id,
                "name": item.name
            })
        return result, 200
    

    @handle_exceptions
    def options_consumption(self):
        consumptions, cc = self.complaint_repository.get_options_consumption() 
        if cc != 200:
            return consumptions, cc
        
        result = []
        for consumption in consumptions:
            result.append({
                "id": consumption.id,
                "name": consumption.name
            })
        return result, 200


    @handle_exceptions
    def options_categories(self):
        categories, cc = self.complaint_repository.get_options_categories() 
        if cc != 200:
            return categories, cc
        
        result = []
        for category in categories:
            result.append({
                "id": category.id,
                "name": category.name
            })
        return result, 200
    

    @handle_exceptions
    def chat(self, data):
        complaint_id = data.get("complaint_id")
        if not complaint_id:
            return "complaint_id requerido", 400

        comment = (data.get("comment") or "").strip()
        if not comment:
            return "Comentario vacío", 400

        current_user_id = int(get_jwt_identity())
        user, uc = self.user_repository.get_user_by_id(current_user_id)
        if uc != 200:
            return user, uc
        
        user_name = user.name

        complaint, cc = self.complaint_repository.get_complaint(complaint_id)
        if cc != 200:
            return complaint, cc

        complaint_letter = "RC" if complaint.complaint_type_id == 1 else "QJ"

        chat, cc = self.complaint_repository.add_chat(
            complaint_id=complaint_id,
            user_id=current_user_id,
            comment=comment
        )
        if cc != 200:
            return chat, cc

        participants, _ = self.complaint_repository.get_chat_participants(
            complaint_id=complaint_id,
            exclude_user_id=current_user_id,
            include_owners=True,
        )
        
        self.push_service.send_to_users(
            user_ids=participants,
            title=f"Nuevo mensaje, reclamo {complaint_letter}-{complaint_id}",
            body=f"{format_name(user_name, True)}: {comment}",
        )
        socketio.emit("complaint_dashboard_update", {})

        return {
            "id": chat.id,
            "comment": chat.comment,
            "commenter_id": chat.commenter_id,
            "commenter_name": format_name(chat.commenter.name),
            "commenter_image": chat.commenter.image,
            "created_at": format_datetime(chat.created_at),
        }, 200
